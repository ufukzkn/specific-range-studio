from __future__ import annotations

import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BENCHMARK_PATH = ROOT / "artifacts" / "benchmarks" / "pi_benchmark_latest.json"
OUT_DIR = ROOT / "docs" / "comparison_reports"
FIG_DIR = OUT_DIR / "figures_updated"
DOCX_PATH = OUT_DIR / "Guncel_Karsilastirma_ve_Benchmark_Grafikleri.docx"

MODEL_LABELS = {
    "interpolation": "İnterpolasyon",
    "xgboost": "XGBoost",
    "ft_transformer": "FT-Transformer",
}
COLORS = {
    "interpolation": "#2f80ed",
    "xgboost": "#22a06b",
    "ft_transformer": "#f2994a",
}


def _load_benchmark() -> dict:
    if not BENCHMARK_PATH.exists():
        raise FileNotFoundError(f"Benchmark JSON not found: {BENCHMARK_PATH}")
    return json.loads(BENCHMARK_PATH.read_text(encoding="utf-8"))


def _model_payload(benchmark: dict, key: str) -> dict:
    payload = (benchmark.get("models") or {}).get(key)
    if not payload:
        raise KeyError(f"Missing benchmark payload for {key}")
    return payload


def _accuracy_component(metrics: dict) -> float:
    targets = {"rmse": 0.0030, "mae": 0.0015, "mape": 2.0, "r2_gap": 0.0010}
    return float(
        np.mean(
            [
                metrics["rmse"] / targets["rmse"],
                metrics["mae"] / targets["mae"],
                metrics["mape"] / targets["mape"],
                max(1.0 - metrics["r2"], 0.0) / targets["r2_gap"],
            ]
        )
    )


def _cost_rows(benchmark: dict) -> dict[str, dict[str, float | None]]:
    weights = {"accuracy": 0.45, "latency": 0.25, "memory": 0.20, "cpu": 0.10}
    ram_budget_mb = 1024.0
    latency_target_ms = 10.0
    cpu_target_percent = 80.0
    rows: dict[str, dict[str, float | None]] = {}
    for key in ("interpolation", "xgboost", "ft_transformer"):
        item = _model_payload(benchmark, key)
        latency_component = item["latency_ms"]["p95"] / latency_target_ms
        memory_component = item["memory_mb"]["rss_peak"] / ram_budget_mb
        cpu_component = item["cpu_percent"]["process_avg"] / cpu_target_percent
        acc = item.get("accuracy") or {}
        if key == "interpolation" or "rmse" not in acc:
            accuracy_component = None
            runtime_cost = (
                (weights["latency"] / 0.55) * latency_component
                + (weights["memory"] / 0.55) * memory_component
                + (weights["cpu"] / 0.55) * cpu_component
            )
            combined = runtime_cost
        else:
            accuracy_component = _accuracy_component(acc)
            combined = (
                weights["accuracy"] * accuracy_component
                + weights["latency"] * latency_component
                + weights["memory"] * memory_component
                + weights["cpu"] * cpu_component
            )
        rows[key] = {
            "accuracy": accuracy_component,
            "latency": latency_component,
            "memory": memory_component,
            "cpu": cpu_component,
            "combined": combined,
            "fit_score": 100.0 / (1.0 + 0.30 * max(combined, 0.0)),
        }
    return rows


def _savefig(path: Path) -> None:
    plt.tight_layout()
    plt.savefig(path, dpi=220, bbox_inches="tight")
    plt.close()


def plot_accuracy(benchmark: dict) -> Path:
    path = FIG_DIR / "01_learned_accuracy_corrected.png"
    models = ["xgboost", "ft_transformer"]
    metrics = ["rmse", "mae", "mape"]
    values = np.array([[_model_payload(benchmark, model)["accuracy"][metric] for metric in metrics] for model in models])
    x = np.arange(len(metrics))
    width = 0.34
    fig, ax = plt.subplots(figsize=(8.4, 4.8))
    for idx, model in enumerate(models):
        ax.bar(x + (idx - 0.5) * width, values[idx], width, label=MODEL_LABELS[model], color=COLORS[model])
    ax.set_xticks(x, ["RMSE", "MAE", "MAPE (%)"])
    ax.set_title("Öğrenilmiş Modellerde Doğruluk Kıyaslaması", fontweight="bold")
    ax.set_ylabel("Hata değeri (düşük daha iyi)")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False)
    ax.text(
        0.0,
        -0.22,
        "Not: İnterpolasyon referans aile olduğu için accuracy yarışına dahil edilmemiştir.",
        transform=ax.transAxes,
        fontsize=9,
        color="#5b677a",
    )
    _savefig(path)
    return path


def plot_latency(benchmark: dict) -> Path:
    path = FIG_DIR / "02_latency_p50_p95.png"
    models = ["interpolation", "xgboost", "ft_transformer"]
    p50 = [_model_payload(benchmark, model)["latency_ms"]["p50"] for model in models]
    p95 = [_model_payload(benchmark, model)["latency_ms"]["p95"] for model in models]
    x = np.arange(len(models))
    width = 0.34
    fig, ax = plt.subplots(figsize=(8.4, 4.8))
    ax.bar(x - width / 2, p50, width, label="p50", color="#80bfff")
    ax.bar(x + width / 2, p95, width, label="p95", color="#1b75d0")
    ax.set_xticks(x, [MODEL_LABELS[m] for m in models])
    ax.set_ylabel("Gecikme (ms)")
    ax.set_title("Tekil Çıkarım Gecikmesi", fontweight="bold")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False)
    _savefig(path)
    return path


def plot_resources(benchmark: dict) -> Path:
    path = FIG_DIR / "03_resource_profile.png"
    models = ["interpolation", "xgboost", "ft_transformer"]
    labels = [MODEL_LABELS[m] for m in models]
    rss = [_model_payload(benchmark, m)["memory_mb"]["rss_peak"] for m in models]
    size = [_model_payload(benchmark, m)["model_size_mb"] for m in models]
    cpu = [_model_payload(benchmark, m)["cpu_percent"]["process_avg"] for m in models]
    fig, axes = plt.subplots(1, 3, figsize=(11.4, 3.7))
    for ax, vals, title, ylabel in [
        (axes[0], rss, "Peak RSS", "MB"),
        (axes[1], cpu, "Ortalama CPU", "%"),
        (axes[2], size, "Model / Artefact Boyutu", "MB"),
    ]:
        ax.bar(labels, vals, color=[COLORS[m] for m in models])
        ax.set_title(title, fontweight="bold", fontsize=11)
        ax.set_ylabel(ylabel)
        ax.tick_params(axis="x", rotation=18)
        ax.grid(axis="y", alpha=0.22)
    fig.suptitle("Kaynak Kullanımı Profili", fontweight="bold", y=1.04)
    _savefig(path)
    return path


def plot_cost_components(benchmark: dict) -> Path:
    path = FIG_DIR / "04_cost_components_corrected.png"
    rows = _cost_rows(benchmark)
    models = ["interpolation", "xgboost", "ft_transformer"]
    categories = ["Accuracy", "Latency", "Memory", "CPU", "Combined"]
    x = np.arange(len(categories))
    width = 0.25
    fig, ax = plt.subplots(figsize=(9.5, 4.9))
    for idx, model in enumerate(models):
        vals = [
            rows[model]["accuracy"] if rows[model]["accuracy"] is not None else 0.0,
            rows[model]["latency"],
            rows[model]["memory"],
            rows[model]["cpu"],
            rows[model]["combined"],
        ]
        bars = ax.bar(x + (idx - 1) * width, vals, width, label=MODEL_LABELS[model], color=COLORS[model])
        if model == "interpolation":
            bars[0].set_hatch("//")
            bars[0].set_alpha(0.30)
    ax.set_xticks(x, categories)
    ax.set_ylabel("Normalize maliyet (düşük daha iyi)")
    ax.set_title("Maliyet Bileşenleri", fontweight="bold")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False)
    ax.text(x[0] - width, 0.04, "N/A", ha="center", va="bottom", fontsize=9, color="#36516f", fontweight="bold")
    ax.text(
        0.0,
        -0.25,
        "Taralı İnterpolasyon/Accuracy sütunu N/A anlamındadır; referans aile accuracy cost hesabına katılmaz.",
        transform=ax.transAxes,
        fontsize=9,
        color="#5b677a",
    )
    _savefig(path)
    return path


def plot_fit_score(benchmark: dict) -> Path:
    path = FIG_DIR / "05_fit_score.png"
    rows = _cost_rows(benchmark)
    models = ["interpolation", "xgboost", "ft_transformer"]
    scores = [rows[m]["fit_score"] for m in models]
    fig, ax = plt.subplots(figsize=(8.4, 4.4))
    ax.bar([MODEL_LABELS[m] for m in models], scores, color=[COLORS[m] for m in models])
    ax.set_ylim(0, max(scores) * 1.2)
    ax.set_ylabel("Fit score (yüksek daha iyi)")
    ax.set_title("Benchmark Tabanlı Uygunluk Skoru", fontweight="bold")
    ax.grid(axis="y", alpha=0.22)
    _savefig(path)
    return path


def plot_radar() -> Path:
    path = FIG_DIR / "06_profile_radar_corrected.png"
    axes_labels = ["Referans rolü", "Hız", "Düşük RAM", "Küçük boyut", "Açıklanabilirlik", "Deployment"]
    scores = {
        "interpolation": [5.0, 3.4, 5.0, 3.5, 5.0, 4.0],
        "xgboost": [4.5, 5.0, 3.6, 4.0, 4.0, 4.6],
        "ft_transformer": [4.0, 4.3, 2.6, 5.0, 2.5, 3.2],
    }
    angles = np.linspace(0, 2 * np.pi, len(axes_labels), endpoint=False).tolist()
    angles += angles[:1]
    fig, ax = plt.subplots(figsize=(8.4, 7.2), subplot_kw={"polar": True})
    for model, vals in scores.items():
        closed = vals + vals[:1]
        ax.plot(angles, closed, label=MODEL_LABELS[model], color=COLORS[model], linewidth=2.2)
        ax.fill(angles, closed, color=COLORS[model], alpha=0.12)
    ax.set_xticks(angles[:-1], axes_labels)
    ax.set_yticks([1, 2, 3, 4, 5])
    ax.set_ylim(0, 5)
    ax.set_title("Genel Profil Radar Grafiği\n(Accuracy Ayrı Değerlendirildi)", fontweight="bold", pad=24, fontsize=15)
    ax.legend(loc="lower center", bbox_to_anchor=(0.5, -0.18), ncol=3, frameon=False)
    fig.text(
        0.10,
        0.01,
        "Not: İnterpolasyonun accuracy skoru yoktur; referans aile rolü ayrı eksende gösterilmiştir.",
        fontsize=9,
        color="#5b677a",
    )
    _savefig(path)
    return path


def _set_doc_styles(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [
        ("Title", 21, RGBColor(21, 55, 100)),
        ("Heading 1", 15, RGBColor(21, 88, 152)),
        ("Heading 2", 12, RGBColor(54, 69, 88)),
    ]:
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = color


def _add_picture(doc: Document, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.4))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)


def build_docx(benchmark: dict, figures: list[Path]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    _set_doc_styles(doc)

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Güncel Karşılaştırma ve Benchmark Grafikleri").bold = True
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("İnterpolasyon referans aile olarak ayrıldı; accuracy kıyası yalnız öğrenilmiş modeller için verildi.").italic = True

    doc.add_heading("Kullanım Notu", level=1)
    doc.add_paragraph(
        "Bu belge, mevcut karşılaştırma ve benchmark raporlarındaki görsel dili koruyarak güncellenmiş grafik setini verir. "
        "Özellikle radar ve accuracy grafiklerinde interpolasyon artık öğrenilmiş model gibi puanlanmaz; benchmark tarafında yalnız runtime, bellek, CPU ve artefact boyutu ile gösterilir."
    )

    run_id = benchmark.get("run_id", "-")
    system = benchmark.get("system", {})
    doc.add_paragraph(
        f"Benchmark kaynağı: {BENCHMARK_PATH.name} | Run ID: {run_id} | Görünen hedef profil: {system.get('display_name', '-')}"
    )

    captions = [
        "Şekil 1. XGBoost ve FT-Transformer için hata metrikleri. İnterpolasyon referans aile olduğu için bu grafikte yoktur.",
        "Şekil 2. Üç yöntemin p50 ve p95 tekil çıkarım gecikmeleri.",
        "Şekil 3. Üç yöntemin bellek, CPU ve artefact/model boyutu profili.",
        "Şekil 4. Normalize maliyet bileşenleri. İnterpolasyon accuracy cost hesabına katılmamıştır.",
        "Şekil 5. Benchmark tabanlı uygunluk skoru. İnterpolasyon için skor runtime-only bağlamında yorumlanmalıdır.",
        "Şekil 6. Genel profil radar grafiği. Accuracy ekseni yerine interpolasyonun referans rolü ayrı gösterilmiştir.",
    ]
    for idx, (figure, caption) in enumerate(zip(figures, captions, strict=True), 1):
        doc.add_heading(f"Grafik {idx}", level=1)
        _add_picture(doc, figure, caption)

    doc.add_heading("Kısa Yorum", level=1)
    doc.add_paragraph(
        "Güncel görseller, mevcut uygulama mantığıyla uyumludur: interpolasyon deterministic referans aile, XGBoost pratik ve güçlü baseline, FT-Transformer ise akademik mimari/araştırma adayıdır. "
        "Bu nedenle accuracy metrikleri XGBoost ve FT-Transformer arasında yorumlanmalı; interpolasyonun doğruluğu aynı eksende düşük veya yüksek puanlanmamalıdır."
    )
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)


def main() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    benchmark = _load_benchmark()
    figures = [
        plot_accuracy(benchmark),
        plot_latency(benchmark),
        plot_resources(benchmark),
        plot_cost_components(benchmark),
        plot_fit_score(benchmark),
        plot_radar(),
    ]
    build_docx(benchmark, figures)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
